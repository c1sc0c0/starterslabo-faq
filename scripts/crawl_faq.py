#!/usr/bin/env python3
"""Log in to mijn.starterslabo.be, crawl FAQ articles + attachments, write a Q&A corpus.

Credentials come from --email/--password or STARTERSLABO_EMAIL / STARTERSLABO_PASSWORD.
Never logs the password. Never writes credentials to disk.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import date
from pathlib import Path

from playwright.sync_api import sync_playwright

LOGIN_URL = "https://mijn.starterslabo.be/login.aspx"
FAQ_URL = "https://mijn.starterslabo.be/Views/FAQ.aspx"
PUBLIC_FAQ_URLS = (
    "https://starterslabo.be/faq/",
    "https://starterslabo.be/faq/page/2/",
)

EXTRACT_ARTICLES = """() => {
  const abs = (href) => {
    try { return new URL(href, location.href).href; } catch { return href; }
  };
  const cards = [...document.querySelectorAll(".card")].filter((c) =>
    c.querySelector(".card-header h2")
  );
  return {
    url: location.href,
    count: cards.length,
    articles: cards.map((card, idx) => {
      const spans = [...card.querySelectorAll(".card-header h2 span")];
      const title = (spans[0]?.textContent || "").trim();
      const meta = (spans[1]?.textContent || "").trim();
      const body = card.querySelector(".card-body");
      const attachBox = [...card.querySelectorAll("a[href*='DynamicDownload']")];
      const allLinks = [...card.querySelectorAll("a[href]")].map((a) => ({
        text: (a.textContent || a.getAttribute("title") || "").trim().replace(/\\s+/g, " "),
        href: abs(a.getAttribute("href") || ""),
      })).filter((l) => l.href && !l.href.includes("javascript:"));
      return {
        index: idx,
        title,
        meta,
        text: body ? body.textContent.trim().replace(/\\n{3,}/g, "\\n\\n") : "",
        attachments: attachBox.map((a) => ({
          text: (a.textContent || "").trim().replace(/\\s+/g, " "),
          href: abs(a.getAttribute("href") || ""),
        })),
        links: allLinks,
      };
    }),
  };
}"""

EXTRACT_PUBLIC = """() => {
  const main =
    document.querySelector("main, .site-main, #content, .entry-content")
    || document.body;
  const blocks = [];
  const headings = [...main.querySelectorAll("h2, h3, h4")];
  for (const h of headings) {
    const title = (h.textContent || "").trim();
    if (!title || title.length < 8) continue;
    if (/vond je dit|recente|archieven|categorie|meta|ontvang onze/i.test(title)) continue;
    const parts = [];
    let n = h.nextElementSibling;
    while (n && !/^H[1-4]$/.test(n.tagName)) {
      const t = (n.textContent || "").trim();
      if (t) parts.push(t);
      n = n.nextElementSibling;
    }
    const text = parts.join("\\n\\n").trim();
    if (text) blocks.push({ title, text });
  }
  return { url: location.href, title: document.title, blocks };
}"""

NAV_SKIP = {
    "Home",
    "Mijn profiel",
    "FAQ",
    "Contact",
    "Uitloggen",
    "Over Starterslabo",
}


def slugify(value: str) -> str:
    value = re.sub(r"[^a-zA-Z0-9._-]+", "-", value).strip("-")
    return value[:120] or "attachment"


def resolve_credentials(email: str | None, password: str | None) -> tuple[str, str]:
    email = (email or os.environ.get("STARTERSLABO_EMAIL") or "").strip()
    password = password or os.environ.get("STARTERSLABO_PASSWORD") or ""
    if not email or not password:
        print(
            "Missing credentials. Pass --email and --password, or set "
            "STARTERSLABO_EMAIL and STARTERSLABO_PASSWORD.",
            file=sys.stderr,
        )
        raise SystemExit(2)
    return email, password


def login(page, email: str, password: str) -> None:
    page.goto(LOGIN_URL, wait_until="domcontentloaded")
    page.fill("#InputUsername", email)
    page.fill("#InputPassword", password)
    page.click("#BtnSubmit")
    page.wait_for_url(lambda url: "login.aspx" not in url.lower(), timeout=30_000)
    if "login.aspx" in page.url.lower():
        raise RuntimeError("Login failed; still on login.aspx")


def unique_path(directory: Path, name: str) -> Path:
    dest = directory / name
    if not dest.exists():
        return dest
    stem, suffix = dest.stem, dest.suffix
    i = 2
    while True:
        candidate = directory / f"{stem}-{i}{suffix}"
        if not candidate.exists():
            return candidate
        i += 1


def download_attachments(page, articles: list[dict], attach_dir: Path) -> dict[str, str]:
    downloaded: dict[str, str] = {}
    attach_dir.mkdir(parents=True, exist_ok=True)
    for article in articles:
        for link in article.get("attachments") or []:
            href = link.get("href") or ""
            name = slugify(link.get("text") or "attachment")
            if not href or href in downloaded:
                continue
            dest = unique_path(attach_dir, name)
            try:
                with page.expect_download(timeout=45_000) as dl_info:
                    page.evaluate(
                        """(url) => {
                          const a = document.createElement('a');
                          a.href = url;
                          a.download = '';
                          document.body.appendChild(a);
                          a.click();
                          a.remove();
                        }""",
                        href,
                    )
                download = dl_info.value
                suggested = download.suggested_filename
                if suggested and suggested.lower() not in {
                    "dynamicdownload.pdf",
                    "dynamicdownload",
                    "download",
                }:
                    dest = unique_path(attach_dir, suggested)
                download.save_as(dest)
                downloaded[href] = dest.name
                print(f"Downloaded {dest.name} ({dest.stat().st_size} bytes)")
            except Exception as exc:
                downloaded[href] = f"ERROR: {exc}"
                print(f"Failed {name}: {exc}", file=sys.stderr)
    return downloaded


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"--- page {i} ---\n{text}")
    return "\n\n".join(parts)


def extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for ti, table in enumerate(doc.tables, 1):
        rows = []
        for row in table.rows:
            cells = [" ".join(c.text.split()) for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            tables.append(f"[table {ti}]\n" + "\n".join(rows))
    return "\n\n".join(paras + tables)


def extract_xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    chunks = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            vals = ["" if v is None else str(v).strip() for v in row]
            if any(vals):
                rows.append("\t".join(vals))
            if i >= 250:
                rows.append("... truncated ...")
                break
        chunks.append(f"## Sheet: {name}\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(chunks)


def extract_attachments(attach_dir: Path, out_dir: Path) -> list[tuple[str, int]]:
    out_dir.mkdir(parents=True, exist_ok=True)
    summaries: list[tuple[str, int]] = []
    if not attach_dir.exists():
        return summaries
    for path in sorted(attach_dir.iterdir()):
        if path.name.startswith("."):
            continue
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf":
                text = extract_pdf(path)
            elif suffix == ".docx":
                text = extract_docx(path)
            elif suffix in {".xlsx", ".xlsm"}:
                text = extract_xlsx(path)
            else:
                continue
        except Exception as exc:
            print(f"Extract failed {path.name}: {exc}", file=sys.stderr)
            continue
        dest = out_dir / f"{path.stem}.txt"
        dest.write_text(text, encoding="utf-8")
        summaries.append((path.name, len(text.strip())))
        print(f"Extracted {path.name}: {len(text.strip())} chars")
    return summaries


def useful_links(article: dict) -> list[dict]:
    extra = []
    seen: set[tuple[str, str]] = set()
    for link in article.get("links") or []:
        text, href = link.get("text") or "", link.get("href") or ""
        if not href or text in NAV_SKIP:
            continue
        if "DynamicDownload" in href:
            continue
        key = (text, href)
        if key in seen:
            continue
        seen.add(key)
        extra.append(link)
    return extra


def write_portal_markdown(articles: list[dict], dest: Path) -> None:
    lines = [
        "# Starters Labo portal FAQ",
        "",
        f"Source: {FAQ_URL}",
        f"Crawled: {date.today().isoformat()}",
        f"Articles: {len(articles)}",
        "",
    ]
    for article in articles:
        title = (article.get("title") or "").lstrip("- ").strip()
        lines.append(f"## {title}")
        lines.append("")
        if article.get("meta"):
            lines.append(f"*{article['meta']}*")
            lines.append("")
        lines.append((article.get("text") or "").strip())
        lines.append("")
        atts = article.get("attachments") or []
        if atts:
            lines.append("**Bijlagen:**")
            for att in atts:
                lines.append(f"- {att.get('text')}")
            lines.append("")
        extra = useful_links(article)
        if extra:
            lines.append("**Links:**")
            for link in extra:
                lines.append(f"- {link['text']}: {link['href']}")
            lines.append("")
        lines.append("")
    dest.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def write_public_markdown(pages: list[dict], dest: Path) -> None:
    lines = [
        "# Public FAQ — starterslabo.be/faq",
        "",
        f"Crawled: {date.today().isoformat()}",
        "",
    ]
    seen: set[str] = set()
    for page in pages:
        for block in page.get("blocks") or []:
            title = (block.get("title") or "").strip()
            if title in seen:
                continue
            seen.add(title)
            lines.append(f"## {title}")
            lines.append("")
            lines.append((block.get("text") or "").strip())
            lines.append("")
    dest.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def crawl_public(page) -> list[dict]:
    pages = []
    for url in PUBLIC_FAQ_URLS:
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=30_000)
            page.wait_for_timeout(800)
            pages.append(page.evaluate(EXTRACT_PUBLIC))
            print(f"Public FAQ {url}: {len(pages[-1].get('blocks') or [])} blocks")
        except Exception as exc:
            print(f"Public FAQ failed {url}: {exc}", file=sys.stderr)
    return pages


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Crawl mijn.starterslabo.be FAQ into a local markdown corpus."
    )
    parser.add_argument("--email", help="Portal username/email (or STARTERSLABO_EMAIL)")
    parser.add_argument("--password", help="Portal password (or STARTERSLABO_PASSWORD)")
    parser.add_argument(
        "--output",
        default="starterslabo-faq-corpus",
        help="Output directory (default: ./starterslabo-faq-corpus)",
    )
    parser.add_argument("--headed", action="store_true", help="Show the browser")
    parser.add_argument(
        "--skip-downloads",
        action="store_true",
        help="Do not download FAQ attachments",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    email, password = resolve_credentials(args.email, args.password)
    output = Path(args.output).expanduser().resolve()
    attach_dir = output / "attachments"
    extracted_dir = output / "extracted"
    output.mkdir(parents=True, exist_ok=True)

    headed = args.headed or os.environ.get("HEADED", "").lower() in ("1", "true", "yes")

    with sync_playwright() as playwright:
        try:
            browser = playwright.chromium.launch(headless=not headed)
        except Exception as exc:
            print(
                "Playwright Chromium is missing. Run: playwright install chromium\n"
                f"Original error: {exc}",
                file=sys.stderr,
            )
            return 1
        page = browser.new_page()
        try:
            login(page, email, password)
            print("Login OK")
            page.goto(FAQ_URL, wait_until="domcontentloaded", timeout=60_000)
            page.wait_for_timeout(1200)
            data = page.evaluate(EXTRACT_ARTICLES)
            print(f"Extracted {data['count']} portal FAQ articles")
            for article in data["articles"]:
                print(
                    f"  [{article['index']}] {article['title'][:70]} "
                    f"({len(article['text'])} chars, {len(article['attachments'])} files)"
                )

            downloads: dict[str, str] = {}
            if not args.skip_downloads:
                downloads = download_attachments(page, data["articles"], attach_dir)

            public_pages = crawl_public(page)
        finally:
            browser.close()

    (output / "manifest.json").write_text(
        json.dumps(
            {
                "url": FAQ_URL,
                "crawled": date.today().isoformat(),
                "article_count": data["count"],
                "articles": data["articles"],
                "downloads": downloads,
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    write_portal_markdown(data["articles"], output / "FAQ.md")
    write_public_markdown(public_pages, output / "public-faq.md")
    if not args.skip_downloads:
        extract_attachments(attach_dir, extracted_dir)
    print(f"Wrote corpus to {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
